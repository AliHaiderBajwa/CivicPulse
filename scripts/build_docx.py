#!/usr/bin/env python3
"""Regenerate docs/CivicPulse-Documentation.docx from the markdown sources.

Run after every milestone (see AGENTS.md):  python3 scripts/build_docx.py
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "CivicPulse-Documentation.docx"

SOURCES = [
    ROOT / "PROGRESS.md",
    ROOT / "docs" / "ASSIGNMENT-BRIEF.md",
    ROOT / "docs" / "RUBRIC-CHECKLIST.md",
    ROOT / "docs" / "PLAN.md",
    ROOT / "docs" / "RUNBOOK.md",
    ROOT / "docs" / "TRIAGE.md",
    ROOT / "docs" / "ENGINEERING-NOTES.md",
    ROOT / "docs" / "AI-USAGE.md",
    ROOT / "README.md",
]
if (ROOT / "docs" / "adr").exists():
    SOURCES.extend(sorted((ROOT / "docs" / "adr").glob("*.md")))

BOLD = re.compile(r"\*\*(.+?)\*\*")
INLINE_CODE = re.compile(r"`([^`]+)`")


def add_inline(paragraph, text: str) -> None:
    """Render **bold** and `code` inline styles."""
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*|`([^`]+)`", text):
        if m.start() > pos:
            paragraph.add_run(text[pos : m.start()])
        if m.group(1) is not None:
            run = paragraph.add_run(m.group(1))
            run.bold = True
        else:
            run = paragraph.add_run(m.group(2))
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0xA3, 0x15, 0x15)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def is_table_line(line: str) -> bool:
    s = line.strip()
    return s.startswith("|") and s.endswith("|")


def render_table(doc: Document, lines: list[str]) -> None:
    rows = [[c.strip() for c in line.strip().strip("|").split("|")] for line in lines]
    rows = [r for r in rows if not all(set(c) <= set("-: ") for c in r)]
    if not rows:
        return
    table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            target = table.cell(i, j).paragraphs[0]
            add_inline(target, cell.replace("**", ""))
            if i == 0:
                for run in target.runs:
                    run.bold = True


def render(doc: Document, text: str) -> None:
    lines = text.splitlines()
    i, in_code, in_table = 0, False, False
    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            if in_code:
                in_code = False
            else:
                in_code = True
                doc.add_paragraph("")
            i += 1
            continue
        if in_code:
            run = doc.add_paragraph().add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            i += 1
            continue

        if is_table_line(line):
            block = []
            while i < len(lines) and is_table_line(lines[i]):
                block.append(lines[i])
                i += 1
            render_table(doc, block)
            continue

        s = line.strip()
        if not s or s == "---":
            i += 1
            continue
        if s.startswith("<!--"):
            i += 1
            continue

        m = re.match(r"(#{1,4})\s+(.*)", s)
        if m:
            level = len(m.group(1))
            if level == 1:
                h = doc.add_heading(m.group(2).replace("**", ""), level=0)
            else:
                h = doc.add_heading("", level=level - 1)
                add_inline(h, m.group(2))
            i += 1
            continue

        m = re.match(r"- \[( |x)\]\s+(.*)", s, re.I)
        if m:
            mark = "☑" if m.group(1).lower() == "x" else "☐"
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, f"{mark} {m.group(2)}")
            i += 1
            continue

        m = re.match(r"^[-*]\s+(.*)", s)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(1))
            i += 1
            continue

        m = re.match(r"^\d+\.\s+(.*)", s)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(1))
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline(p, s)
        i += 1


def main() -> int:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    title = doc.add_heading("CivicPulse — Project Documentation", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("CS4032 Software Construction and Design · Assignment 1 · Auto-generated from markdown — do not hand-edit").italic = True

    written = 0
    for src in SOURCES:
        if not src.exists():
            continue
        doc.add_page_break()
        doc.add_heading(f"Source: {src.relative_to(ROOT)}", level=1)
        render(doc, src.read_text(encoding="utf-8"))
        written += 1

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)} from {written} sources")
    return 0


if __name__ == "__main__":
    sys.exit(main())
